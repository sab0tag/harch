"""XLSX exports (openpyxl). Numbers are written as numbers (kg with 3 decimals) so accountants can sum."""
import calendar
from collections import OrderedDict
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

KG = Decimal(1000)
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EFEA")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
KG_FMT = "#,##0.000"
MONEY_FMT = "#,##0.00"


def _kg(grams):
    return float((Decimal(grams) / KG).quantize(Decimal("0.001")))


def _response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _header(ws, row, titles, widths=None):
    for col, title in enumerate(titles, 1):
        c = ws.cell(row=row, column=col, value=title)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def consumption_xlsx(data, d_from, d_to, group):
    from apps.consumption.services import GROUPS

    wb = Workbook()
    ws = wb.active
    ws.title = "Розхід"
    period = d_from.strftime("%d.%m.%Y") + ("" if d_from == d_to else " – " + d_to.strftime("%d.%m.%Y"))
    ws["A1"] = settings.HOSPITAL_NAME
    ws["A2"] = f"Розхід продуктів {period} ({GROUPS[group]})"
    ws["A2"].font = Font(bold=True, size=13)
    ws["A3"] = "Джерело: зафіксований розхід затверджених планів. Сума — за поточними цінами довідника."
    _header(ws, 5, ["Продукт", "Брутто, кг", "Нетто, кг", "Ціна за кг, грн", "Сума, грн"], [36, 14, 14, 16, 16])
    r = 6
    for g in data["groups"]:
        if g["label"]:
            ws.cell(row=r, column=1, value=g["label"]).font = BOLD
            r += 1
        for row in g["rows"]:
            ws.cell(row=r, column=1, value=row["product_name"])
            ws.cell(row=r, column=2, value=_kg(row["brutto_g"])).number_format = KG_FMT
            ws.cell(row=r, column=3, value=_kg(row["netto_g"])).number_format = KG_FMT
            if row["price_per_kg"] is not None:
                ws.cell(row=r, column=4, value=float(row["price_per_kg"])).number_format = MONEY_FMT
                ws.cell(row=r, column=5, value=float(row["cost"])).number_format = MONEY_FMT
            r += 1
        if g["label"]:
            ws.cell(row=r, column=1, value=f"Разом: {g['label']}").font = BOLD
            ws.cell(row=r, column=2, value=_kg(g["brutto_g"])).number_format = KG_FMT
            ws.cell(row=r, column=5, value=float(g["cost"])).number_format = MONEY_FMT
            r += 1
    ws.cell(row=r, column=1, value="РАЗОМ").font = BOLD
    ws.cell(row=r, column=2, value=_kg(data["brutto_g"])).number_format = KG_FMT
    c = ws.cell(row=r, column=5, value=float(data["cost"]))
    c.number_format, c.font = MONEY_FMT, BOLD
    ws.freeze_panes = "A6"
    return _response(wb, f"rozkhid_{d_from.isoformat()}_{d_to.isoformat()}.xlsx")


def accumulative_data(year, month):
    """Product × day of month, kg, from APPROVED write-offs (the accounting source)."""
    from apps.consumption.models import WriteOffLine

    days_in_month = calendar.monthrange(year, month)[1]
    lines = WriteOffLine.objects.filter(
        write_off__approved_at__isnull=False,
        write_off__plan__service_date__year=year, write_off__plan__service_date__month=month,
    ).values_list("product_id", "product_name", "write_off__plan__service_date", "quantity_g", "cost")
    products, closed_days = OrderedDict(), set()
    for pid, name, day, qty, cost in sorted(lines, key=lambda l: (l[1], l[2])):
        p = products.setdefault(pid, {"name": name, "days": {}, "total_g": Decimal(0), "cost": Decimal(0)})
        p["days"][day.day] = p["days"].get(day.day, Decimal(0)) + qty
        p["total_g"] += qty
        p["cost"] += cost
        closed_days.add(day.day)
    return {"days": days_in_month, "products": list(products.values()), "closed_days": closed_days,
            "cost": sum((p["cost"] for p in products.values()), Decimal(0))}


MONTHS = ["", "січень", "лютий", "березень", "квітень", "травень", "червень", "липень", "серпень",
          "вересень", "жовтень", "листопад", "грудень"]


def accumulative_xlsx(year, month):
    data = accumulative_data(year, month)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{month:02d}.{year}"
    ws["A1"] = settings.HOSPITAL_NAME
    ws["A2"] = f"Накопичувальна відомість витрачання продуктів харчування за {MONTHS[month]} {year} р., кг"
    ws["A2"].font = Font(bold=True, size=13)
    open_days = [d for d in range(1, data["days"] + 1) if d not in data["closed_days"]]
    ws["A3"] = ("Дні без затвердженого списання: " + ", ".join(map(str, open_days))) if open_days else \
        "Усі дні місяця мають затверджене списання."
    titles = ["Продукт"] + [str(d) for d in range(1, data["days"] + 1)] + ["Разом, кг", "Сума, грн"]
    _header(ws, 5, titles, [30] + [7] * data["days"] + [12, 14])
    r = 6
    for p in data["products"]:
        ws.cell(row=r, column=1, value=p["name"]).border = BOX
        for d in range(1, data["days"] + 1):
            c = ws.cell(row=r, column=1 + d, value=_kg(p["days"][d]) if d in p["days"] else None)
            c.number_format, c.border = KG_FMT, BOX
        c = ws.cell(row=r, column=data["days"] + 2, value=_kg(p["total_g"]))
        c.number_format, c.font, c.border = KG_FMT, BOLD, BOX
        c = ws.cell(row=r, column=data["days"] + 3, value=float(p["cost"]))
        c.number_format, c.border = MONEY_FMT, BOX
        r += 1
    ws.cell(row=r, column=1, value="РАЗОМ").font = BOLD
    c = ws.cell(row=r, column=data["days"] + 3, value=float(data["cost"]))
    c.number_format, c.font = MONEY_FMT, BOLD
    r += 3
    for label in ("Бухгалтер", "Лікар-дієтолог", "Завідувач виробництва"):
        ws.cell(row=r, column=1, value=f"{label}  ____________________")
        r += 2
    ws.freeze_panes = "B6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    return _response(wb, f"nakopychuvalna_{year}-{month:02d}.xlsx")
