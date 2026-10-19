"""S4-04/S4-05: printable menu-rozkladka, typical forms З-4а / З-13, and the full day end to end."""
import calendar
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from apps.core.testing import on_the_day
from openpyxl import load_workbook

from apps.catalog.models import Product
from apps.consumption.models import WriteOff
from apps.consumption.tests.conftest import DAY, ORDERS, as_user, cooking_day, diet, write_off_day, world  # noqa: F401
from apps.reports.pdf import menu_rozkladka_context
from apps.workflow.definitions import COOKING, MENU_PLANNING, ORDER_SUBMISSION, WRITE_OFF
from apps.workflow.models import Notification, Task, TaskStatus

pytestmark = pytest.mark.django_db


def test_menu_pdf_for_approved_plan(client, cooking_day):
    r = as_user(client, cooking_day.chef).get(f"/reports/menu/{DAY.isoformat()}.pdf")
    assert r["Content-Type"] == "application/pdf" and r.content[:4] == b"%PDF"


def test_menu_pdf_context_matches_plan(cooking_day):
    ctx = menu_rozkladka_context(cooking_day.plan)
    assert {d["code"]: d["portions"] for d in ctx["diets"]} == {"ОВД": 32, "ЩД": 20, "№9": 8}
    assert ctx["portions"] == 60 and all(d["slots"] for d in ctx["diets"])


@pytest.fixture
def requisites(settings):
    settings.HOSPITAL_NAME = "КНП ОМР «ОХТИРСЬКА ЦРЛ»"
    settings.HOSPITAL_FULL_NAME = "КОМУНАЛЬНЕ НЕКОМЕРЦІЙНЕ ПІДПРИЄМСТВО ОХТИРСЬКОЇ МІСЬКОЇ РАДИ «ОХТИРСЬКА ЦЕНТРАЛЬНА РАЙОННА ЛІКАРНЯ»"
    settings.HOSPITAL_EDRPOU = "02007472"
    settings.HOSPITAL_ADDRESS = "42700, м. Охтирка, вул. Петропавлівська, 15"
    settings.APPROVER_TITLE = "Директор"
    return settings


def test_printed_menu_carries_official_requisites(cooking_day, requisites):
    from django.template.loader import render_to_string

    html = render_to_string("reports/menu_rozkladka.html", menu_rozkladka_context(cooking_day.plan))
    assert requisites.HOSPITAL_FULL_NAME in html and "ЄДРПОУ 02007472" in html and "Петропавлівська, 15" in html
    assert "ЗАТВЕРДЖУЮ<br>Директор" in html and "Головний лікар" not in html


def test_typical_forms_carry_official_requisites(client, cooking_day, requisites):
    ws = _vymoha(client, cooking_day.accountant)
    assert ws["A1"].value == requisites.HOSPITAL_FULL_NAME
    assert ws["A2"].value == "Ідентифікаційний код за ЄДРПОУ 02007472"
    assert any(c.value == "Директор ____________" for c in ws[3])


def test_menu_pdf_refused_before_approval(client, world):
    r = as_user(client, world.dietician).get(f"/reports/menu/{DAY.isoformat()}.pdf")
    assert r.status_code == 302 and "/plan/" in r.url


def test_menu_pdf_access(client, cooking_day):
    assert as_user(client, cooking_day.nurse["Терапевтичне"]).get(f"/reports/menu/{DAY.isoformat()}.pdf").status_code == 403


def test_accumulative_statement_from_approved_write_offs(client, write_off_day):
    w = write_off_day
    Product.objects.update(price_per_kg=Decimal("20.00"))
    as_user(client, w.accountant).post(f"/write-off/{DAY.isoformat()}/", {"action": "submit"})
    month = DAY.strftime("%Y-%m")
    c = as_user(client, w.accountant)
    assert "ще немає затверджених" in c.get(f"/reports/accumulative/{month}/").content.decode()  # submitted ≠ approved
    as_user(client, w.dietician).post(f"/tasks/{w.write_off.pk}/approve/")
    r = as_user(client, w.accountant).get(f"/reports/accumulative/{month}/?format=xlsx")
    ws = load_workbook(BytesIO(r.content)).active
    wo = WriteOff.objects.get(plan=w.plan)
    line = wo.lines.order_by("product_name").first()
    assert ws["A1"].value and any(c.value == "Типова форма № З-13" for c in ws[1])
    row = next(i for i in range(7, ws.max_row + 1) if ws.cell(row=i, column=2).value == line.product_name)
    assert ws.cell(row=row, column=3).value == "кг"
    assert ws.cell(row=row, column=3 + DAY.day).value == pytest.approx(float(line.quantity_g / 1000), abs=0.001)
    assert ws.cell(row=row, column=3 + DAY.day + 1).value is None           # next day has no write-off
    days = calendar.monthrange(DAY.year, DAY.month)[1]
    assert ws.cell(row=row, column=5 + days).value == 20.0                  # average price
    assert ws.cell(row=row, column=6 + days).value == pytest.approx(float(line.cost))


def test_reports_access(client, world):
    assert as_user(client, world.chef).get("/reports/").status_code == 403
    assert as_user(client, world.accountant).get("/reports/").status_code == 200


@on_the_day()
def test_full_day_end_to_end(client, world):
    """Orders from 3 departments → plan → chef → kitchen fact → write-off closed. All through HTTP."""
    day = DAY.isoformat()
    for name, portions in ORDERS.items():
        c = as_user(client, world.nurse[name])
        url = f"/orders/{world.depts[name].pk}/{day}/"
        c.get(url)
        c.post(url, {"action": "submit", **{f"diet_{diet(k).pk}": v for k, v in portions.items()}})
    c = as_user(client, world.dietician)
    c.post("/orders/summary/", {"date": day, "action": "approve_all"})
    plan_task = Task.objects.get(type=MENU_PLANNING, service_date=DAY)
    html = c.get(f"/plan/{day}/").content.decode()
    assert "Подати шеф-кухарю" in html
    from apps.planning.models import DailyPlan

    plan = DailyPlan.objects.get(service_date=DAY)
    c.post(f"/plan/{day}/", {"action": "submit", **{f"menu_{e.diet_id}": e.menu_id for e in plan.entries.all()}})
    as_user(client, world.chef).post(f"/tasks/{plan_task.pk}/approve/")

    chef = as_user(client, world.chef)
    chef.get(f"/kitchen/{day}/")
    fact = {}
    for a in plan.actuals.all():
        fact[f"fact_{a.pk}"] = f"{a.actual_brutto_g / 1000:.3f}"
    chef.post(f"/kitchen/{day}/", {**fact, "action": "submit"})
    cooking = Task.objects.get(type=COOKING, service_date=DAY)
    as_user(client, world.dietician).post(f"/tasks/{cooking.pk}/approve/")

    Product.objects.update(price_per_kg=Decimal("55.00"))
    as_user(client, world.accountant).post(f"/write-off/{day}/", {"action": "submit"})
    write_off = Task.objects.get(type=WRITE_OFF, service_date=DAY)
    as_user(client, world.dietician).post(f"/tasks/{write_off.pk}/approve/")

    statuses = dict(Task.objects.filter(service_date=DAY).values_list("type", "status").distinct())
    assert set(Task.objects.filter(service_date=DAY).values_list("status", flat=True)) == {"approved"}
    assert statuses.keys() == {ORDER_SUBMISSION, MENU_PLANNING, COOKING, WRITE_OFF}
    wo = WriteOff.objects.get(plan=plan)
    assert wo.approved_at and wo.total_cost == sum(l.cost for l in wo.lines.all()) > 0
    # the accountant was told the write-off was approved
    assert Notification.objects.filter(user=world.accountant, task=write_off, title__startswith="Затверджено").exists()


# ---------------------------------------------------------------- S4-05 меню-вимога З-4а
def _vymoha(client, user):
    r = as_user(client, user).get(f"/reports/menu-vymoha/{DAY.isoformat()}.xlsx")
    assert r.status_code == 200, r.status_code
    return load_workbook(BytesIO(r.content)).active


def _vymoha_rows(ws):
    """{product: [values…]} for the table body (header is two rows at 10–11)."""
    rows = {}
    for r in range(12, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if name == "РАЗОМ":
            break
        rows[name] = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    return rows


def test_menu_vymoha_from_plan_before_write_off(client, cooking_day):
    """Eaters by diet and per-meal quantities come from the frozen snapshot; slots add up to the total."""
    from apps.consumption.services import snapshot_by_product

    ws = _vymoha(client, cooking_day.accountant)
    assert any(c.value == "Типова форма № З-4а" for c in ws[1])
    assert "Попередньо" in ws["A6"].value
    eaters = " ".join(str(c.value) for c in ws[8] if c.value)
    assert "ОВД: 32" in eaters and "ЩД: 20" in eaters and "№9: 8" in eaters and "Разом: 60" in eaters
    slot_cols = [c for c in range(4, ws.max_column + 1) if ws.cell(row=11, column=c).value]
    total_col = slot_cols[-1] + 1
    assert ws.cell(row=10, column=total_col).value == "Разом до закладки"
    rows = _vymoha_rows(ws)
    planned = {name: g for name, g in snapshot_by_product(cooking_day.plan).values()}
    assert set(rows) == set(planned)
    for name, values in rows.items():
        by_slot = sum(values[c - 1] or 0 for c in slot_cols)
        assert values[total_col - 1] == pytest.approx(float(planned[name] / 1000), abs=0.002)
        assert by_slot == pytest.approx(values[total_col - 1], abs=0.005)


def test_menu_vymoha_uses_write_off_quantities_and_prices(client, write_off_day):
    w = write_off_day
    Product.objects.update(price_per_kg=Decimal("20.00"))
    a = w.plan.actuals.order_by("product_name").first()
    a.actual_brutto_g = a.planned_brutto_g * 2
    a.reason = "брак"
    a.save()
    as_user(client, w.accountant).post(f"/write-off/{DAY.isoformat()}/", {"action": "submit"})
    Product.objects.update(price_per_kg=Decimal("99.00"))      # later price change must not leak in
    ws = _vymoha(client, w.accountant)
    assert "поданого на перевірку" in ws["A6"].value
    line = WriteOff.objects.get(plan=w.plan).lines.get(product=a.product)
    values = _vymoha_rows(ws)[line.product_name]
    planned, issued, price, cost = values[-4:]
    assert issued == pytest.approx(float(line.quantity_g / 1000), abs=0.001)
    assert issued == pytest.approx(2 * planned, abs=0.002)                   # fact ≠ plan is visible
    assert price == 20.0 and cost == pytest.approx(float(line.cost))
    as_user(client, w.dietician).post(f"/tasks/{w.write_off.pk}/approve/")
    assert "затвердженого списання" in _vymoha(client, w.accountant)["A6"].value


def test_menu_vymoha_refused_before_approval(client, world):
    r = as_user(client, world.accountant).get(f"/reports/menu-vymoha/{DAY.isoformat()}.xlsx")
    assert r.status_code == 302 and r.url == "/reports/"


def test_menu_vymoha_access(client, cooking_day):
    url = f"/reports/menu-vymoha/{DAY.isoformat()}.xlsx"
    from apps.accounts.models import Role
    from apps.accounts.tests.factories import UserFactory

    # the chef signs «Продукти одержав (кухар)» and the store «видав (комірник)» (S7-02)
    assert as_user(client, cooking_day.chef).get(url).status_code == 200
    assert as_user(client, UserFactory(role=Role.STOREKEEPER)).get(url).status_code == 200
    assert as_user(client, cooking_day.nurse["Терапевтичне"]).get(url).status_code == 403
    assert as_user(client, cooking_day.dietician).get(url).status_code == 200


@pytest.mark.parametrize("transition", ["reject", "recall"])
def test_returned_write_off_export_uses_current_draft_prices(write_off_day, transition):
    """A returned draft must not keep exporting the values frozen by its previous submission."""
    from apps.consumption import services
    from apps.consumption.models import WriteOff
    from apps.planning.models import DailyPlan
    from apps.reports.exports import menu_vymoha_data
    from apps.workflow import services as wf

    w = write_off_day
    Product.objects.update(price_per_kg=Decimal("20.00"))
    wf.submit(w.write_off, w.accountant)
    frozen = WriteOff.objects.get(plan=w.plan).lines.first()
    assert frozen.price_per_kg == Decimal("20.00")

    if transition == "reject":
        wf.reject(w.write_off, w.dietician, "Уточніть ціну")
        expected_status = TaskStatus.REJECTED
    else:
        wf.recall(w.write_off, w.accountant)
        expected_status = TaskStatus.IN_PROGRESS

    wo = WriteOff.objects.select_related("task", "plan").get(plan=w.plan)
    assert wo.task.status == expected_status
    services.set_prices(wo, w.accountant, {frozen.product_id: Decimal("30.00")})
    live = next(row for row in services.write_off_rows(wo) if row["product_id"] == frozen.product_id)
    assert live["price_per_kg"] == Decimal("30.00")

    exported = menu_vymoha_data(DailyPlan.objects.get(pk=w.plan.pk))
    row = next(row for row in exported["rows"] if row["name"] == frozen.product_name)
    assert row["price_per_kg"] == live["price_per_kg"]
    assert row["cost"] == live["cost"]
    assert exported["source"] == "fact"

    wf.submit(wo.task, w.accountant)
    resubmitted = WriteOff.objects.get(plan=w.plan)
    assert resubmitted.lines.get(product_id=frozen.product_id).price_per_kg == Decimal("30.00")
    exported = menu_vymoha_data(DailyPlan.objects.get(pk=w.plan.pk))
    row = next(row for row in exported["rows"] if row["name"] == frozen.product_name)
    assert row["price_per_kg"] == Decimal("30.00")
    assert exported["source"] == "submitted"


# ------------------------------------------------------------------ S7-03 бракеражний журнал
def test_brakerazh_pdf_for_the_kitchen(client, cooking_day):
    r = as_user(client, cooking_day.chef).get(f"/reports/brakerazh/{DAY.isoformat()}.pdf")
    assert r["Content-Type"] == "application/pdf" and r.content[:4] == b"%PDF"


def test_brakerazh_lists_every_cooked_dish_with_corrections(cooking_day):
    from apps.consumption.services import kitchen_board
    from apps.orders import corrections
    from apps.reports.pdf import brakerazh_context
    from apps.workflow import services as wf

    w = cooking_day
    ctx = brakerazh_context(w.plan)
    dishes = [d for s in ctx["slots"] for d in s["dishes"]]
    board = [d for s in kitchen_board(w.plan) for d in s["dishes"]]
    assert len(dishes) == len(board) and all(d["kcal"] is not None for d in dishes)
    assert [d["n"] for d in dishes] == list(range(1, len(dishes) + 1))
    before = {(s["label"], d["card_no"]): d["portions"] for s in ctx["slots"] for d in s["dishes"] if "ОВД" in d["diets"]}
    c = corrections.open_correction(w.depts["Терапевтичне"], DAY, w.nurse["Терапевтичне"])
    corrections.submit_correction(c, w.nurse["Терапевтичне"], {diet("ОВД").pk: 3}, from_meal="breakfast",
                                  reason="admission")
    wf.approve(c.task, w.dietician)
    after = {(s["label"], d["card_no"]): d["portions"] for s in brakerazh_context(w.plan)["slots"] for d in s["dishes"]}
    diff = {k: (v, after.get(k)) for k, v in before.items() if after.get(k) != v + 3}
    assert not diff, diff   # the commission sees what was cooked


def test_brakerazh_access_and_before_approval(client, world, cooking_day):
    url = f"/reports/brakerazh/{DAY.isoformat()}.pdf"
    assert as_user(client, cooking_day.nurse["Терапевтичне"]).get(url).status_code == 403
    assert as_user(client, cooking_day.accountant).get(url).status_code == 403
    r = as_user(client, cooking_day.chef).get("/reports/brakerazh/2026-10-02.pdf")
    assert r.status_code == 302 and "/kitchen/" in r.url
