"""S5-07: import product prices from the accounting department's Excel file.

Upload → parsed cells are stored on PriceImport (the file is not kept) → columns are detected or
chosen by hand → each row is matched to a Product → a person ticks what to apply → prices change
in one transaction, with history. Nothing is written without that confirmation. No AI: matching is
normalised names, remembered aliases and difflib.
"""
import difflib
import math
import re
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from .audit import normalize
from .models import PriceImport, PriceImportRow, Product, ProductAlias

MAX_BYTES = 5 * 1024 * 1024
MAX_SHEETS = 10
MAX_ROWS = 3000
MAX_COLS = 30
PREVIEW_ROWS = 8
HEADER_SCAN_ROWS = 15
FUZZY_MIN = Decimal("0.65")
BIG_CHANGE_PCT = 30

NAME_HEADERS = ("найменування", "назва", "продукт", "товар", "наименование", "название")
PRICE_HEADERS = ("ціна", "цiна", "вартість", "цена", "стоимость")
UNIT_HEADERS = ("од. вим", "од.вим", "одиниця", "од вим", "ед. изм", "ед.изм", "единица")
TOTAL_WORDS = ("разом", "всього", "усього", "итого", "всего")

# Price per unit written in the file → multiplier to get UAH per kg. None = cannot convert.
UNIT_FACTORS = {
    "кг": Decimal(1), "kg": Decimal(1), "л": Decimal(1), "l": Decimal(1), "літр": Decimal(1),
    "г": Decimal(1000), "гр": Decimal(1000), "g": Decimal(1000), "мл": Decimal(1000),
    "100 г": Decimal(10), "100г": Decimal(10), "100 гр": Decimal(10),
}
DEFAULT_FACTORS = {PriceImport.Unit.KG: Decimal(1), PriceImport.Unit.G100: Decimal(10), PriceImport.Unit.G: Decimal(1000)}


# ------------------------------------------------------------------ reading
def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")
    return str(value).strip()


def read_xlsx(file) -> dict:
    """{sheet name: rows of cell texts}. Raises ValidationError with a user-facing message."""
    from openpyxl import load_workbook

    name = getattr(file, "name", "") or ""
    if name.lower().endswith(".xls"):
        raise ValidationError("Формат .xls (Excel 97–2003) не підтримується. Відкрийте файл в Excel і збережіть як .xlsx.")
    if getattr(file, "size", 0) > MAX_BYTES:
        raise ValidationError("Файл більший за 5 МБ.")
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — any unreadable upload gets the same message
        raise ValidationError("Не вдалося прочитати файл. Потрібен Excel у форматі .xlsx.") from exc
    sheets = {}
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows = []
            for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
                rows.append([_cell_text(v) for v in row])
            while rows and not any(rows[-1]):
                rows.pop()
            if rows:
                width = max(len(r) for r in rows)
                sheets[ws.title] = [r + [""] * (width - len(r)) for r in rows]
    finally:
        wb.close()
    if not sheets:
        raise ValidationError("У файлі немає даних.")
    return sheets


# ------------------------------------------------------------------ columns
def _find(cells, words):
    for i, c in enumerate(cells):
        low = c.lower()
        if any(w in low for w in words):
            return i
    return None


def detect_mapping(rows):
    """(header_row, name_col, price_col, unit_col, unit) from header titles; Nones when not found."""
    for r, cells in enumerate(rows[:HEADER_SCAN_ROWS]):
        name_col, price_col = _find(cells, NAME_HEADERS), _find(cells, PRICE_HEADERS)
        if name_col is not None and price_col is not None and name_col != price_col:
            unit_col = _find(cells, UNIT_HEADERS)
            title = cells[price_col].lower().replace(" ", "")
            unit = PriceImport.Unit.G100 if "100г" in title else PriceImport.Unit.KG
            return r, name_col, price_col, unit_col, unit
    return None, None, None, None, PriceImport.Unit.KG


def default_mapping(imp: PriceImport):
    """Last applied import's mapping when this file has the same headers there; otherwise detect."""
    prev = PriceImport.objects.filter(status=PriceImport.Status.APPLIED).exclude(pk=imp.pk).first()
    if prev and prev.is_mapped and prev.sheet in imp.sheets and prev.sheet in prev.sheets:
        old, new = prev.sheets[prev.sheet], imp.sheets[prev.sheet]
        if len(old) > prev.header_row and len(new) > prev.header_row and old[prev.header_row] == new[prev.header_row]:
            return prev.sheet, prev.header_row, prev.name_col, prev.price_col, prev.unit_col, prev.unit
    for sheet, rows in imp.sheets.items():
        header_row, name_col, price_col, unit_col, unit = detect_mapping(rows)
        if header_row is not None:
            return sheet, header_row, name_col, price_col, unit_col, unit
    return next(iter(imp.sheets)), None, None, None, None, PriceImport.Unit.KG


def column_letter(i):
    s = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


# ------------------------------------------------------------------ values
def parse_price(text):
    """'45,50' / '1 234.00 грн' / '45.5' → Decimal; None when there is no positive number."""
    t = (text or "").replace("\xa0", " ").lower()
    t = re.sub(r"(грн|uah|₴)\.?", "", t).strip()
    t = re.sub(r"(?<=\d)[ '](?=\d{3}\b)", "", t)      # thousands separators: 1 234,50
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    t = t.replace(",", ".")
    if not re.fullmatch(r"\d+(\.\d+)?", t):
        return None
    try:
        value = Decimal(t)
    except InvalidOperation:
        return None
    return value if value > 0 else None


def unit_factor(raw_unit, default_unit):
    """Multiplier to UAH/kg; None for units that are not weight or volume (шт, уп, пач...)."""
    u = (raw_unit or "").strip().lower().rstrip(".")
    if not u:
        return DEFAULT_FACTORS[default_unit]
    return UNIT_FACTORS.get(u)


# ------------------------------------------------------------------ matching
def _tokens(norm):
    return [w for w in re.findall(r"[^\W\d_]+", norm) if len(w) >= 2]


def _token_sim(a, b):
    """Same word in another form ('гречана' ~ 'гречка', 'курячі' ~ 'куряче'); different stems score 0."""
    if a == b:
        return 1.0
    if a[:3] != b[:3]:
        return 0.0
    return difflib.SequenceMatcher(None, a, b).ratio()


class Matcher:
    """Matches a name from the accountant's file to a catalogue product.

    Words are compared one by one, weighted by how rare they are in the catalogue: 'крупа' or
    'борошно' say little, 'гречка' or 'буряк' say a lot. Word order and extra words in the file
    ('столовий', 'в/г') do not hurt. Different numbers never match (Молоко 2.5% ≠ Молоко 3.2%).
    """

    def __init__(self, products, aliases=None):
        self.products = [(p, normalize(p.name)) for p in products]
        self.aliases = aliases or {}
        self.tokens = {p.pk: _tokens(n) for p, n in self.products}
        df = {}
        for toks in self.tokens.values():
            for t in set(toks):
                df[t] = df.get(t, 0) + 1
        n = max(len(self.products), 1)
        self.weights = {t: math.log(1 + n / c) for t, c in df.items()}
        self.unknown_weight = math.log(1 + n)

    def _weight(self, t):
        return self.weights.get(t, self.unknown_weight)

    def _cover(self, source, target):
        """How much of `source` (weighted) is found in `target`."""
        total = sum(self._weight(t) for t in source)
        if not total or not target:
            return 0.0
        return sum(self._weight(t) * max(_token_sim(t, u) for u in target) for t in source) / total

    def similarity(self, norm_a, product):
        pn = normalize(product.name)
        da, db = re.findall(r"\d+", norm_a), re.findall(r"\d+", pn)
        if da and db and da != db:
            return 0.0                      # Кефір 1% ≠ Кефір 2.5%
        penalty = 0.9 if da != db else 1.0  # a number on one side only: «Масло вершкове 72,5%» ~ «Масло вершкове»
        ta, tb = _tokens(norm_a), self.tokens.get(product.pk) or _tokens(pn)
        # The catalogue name should be fully present in the file's name; extra words there are fine.
        return penalty * (0.75 * self._cover(tb, ta) + 0.25 * self._cover(ta, tb))

    def match(self, raw_name):
        """(product, match, score)."""
        norm = normalize(raw_name)
        if norm in self.aliases:
            return self.aliases[norm], PriceImportRow.Match.ALIAS, Decimal(1)
        for p, pn in self.products:
            if pn == norm:
                return p, PriceImportRow.Match.EXACT, Decimal(1)
        best, best_score = None, 0.0
        for p, _ in self.products:
            score = self.similarity(norm, p)
            if score > best_score:
                best, best_score = p, score
        score = Decimal(str(round(best_score, 2)))
        if best is not None and score >= FUZZY_MIN:
            return best, PriceImportRow.Match.FUZZY, score
        return None, PriceImportRow.Match.NONE, score


# ------------------------------------------------------------------ services
def create_import(file, user) -> PriceImport:
    sheets = read_xlsx(file)
    imp = PriceImport.objects.create(filename=(getattr(file, "name", "") or "файл.xlsx")[:200], uploaded_by=user, sheets=sheets)
    sheet, header_row, name_col, price_col, unit_col, unit = default_mapping(imp)
    set_mapping(imp, sheet, header_row, name_col, price_col, unit_col, unit)
    return imp


def _check_draft(imp):
    if imp.status != PriceImport.Status.DRAFT:
        raise ValidationError("Цей імпорт уже закрито — завантажте файл ще раз.")


@transaction.atomic
def set_mapping(imp, sheet, header_row, name_col, price_col, unit_col=None, unit=PriceImport.Unit.KG):
    """Save the column choice and rebuild matched rows. Incomplete mapping just clears the rows."""
    _check_draft(imp)
    if sheet not in imp.sheets:
        raise ValidationError("Такого аркуша немає у файлі.")
    rows = imp.sheets[sheet]
    width = len(rows[0]) if rows else 0
    for col in (name_col, price_col, unit_col):
        if col is not None and not 0 <= col < width:
            raise ValidationError("Обраного стовпця немає на аркуші.")
    if header_row is not None and not 0 <= header_row < len(rows):
        raise ValidationError("Обраного рядка заголовків немає на аркуші.")
    if name_col is not None and name_col == price_col:
        raise ValidationError("Назва і ціна мають бути в різних стовпцях.")
    imp.sheet, imp.header_row, imp.name_col, imp.price_col, imp.unit_col = sheet, header_row, name_col, price_col, unit_col
    imp.unit = unit
    imp.save()
    imp.rows.all().delete()
    if imp.is_mapped:
        PriceImportRow.objects.bulk_create(build_rows(imp))
    return imp


def build_rows(imp):
    aliases = {a.name: a.product for a in ProductAlias.objects.select_related("product") if a.product.is_active}
    matcher = Matcher(Product.objects.filter(is_active=True), aliases)
    out = []
    for i, cells in enumerate(imp.sheets[imp.sheet][imp.header_row + 1:], start=imp.header_row + 2):
        name, raw_price = cells[imp.name_col].strip(), cells[imp.price_col].strip()
        if not name or not raw_price:
            continue                                   # section titles, blank lines
        if normalize(name).split(" ")[0] in TOTAL_WORDS:
            continue
        raw_unit = cells[imp.unit_col].strip() if imp.unit_col is not None else ""
        row = PriceImportRow(price_import=imp, row_no=i, raw_name=name[:300], raw_price=raw_price[:50], raw_unit=raw_unit[:30])
        price, factor = parse_price(raw_price), unit_factor(raw_unit, imp.unit)
        if price is None:
            row.problem = "Не вдалося прочитати ціну."
        elif factor is None:
            row.problem = f"Одиниця «{raw_unit}» — не вага й не об'єм, ціну за кг не порахувати."
        else:
            row.price_per_kg = (price * factor).quantize(Decimal("0.01"))
        row.product, row.match, row.score = matcher.match(name)
        row.old_price = row.product.price_per_kg if row.product else None
        # Pre-tick only sure matches with a readable price; similar names wait for a person.
        row.selected = (not row.problem and row.match in (PriceImportRow.Match.ALIAS, PriceImportRow.Match.EXACT)
                        and row.price_per_kg != row.old_price)
        out.append(row)
    return out


@transaction.atomic
def apply(imp, user, choices):
    """choices: {row_id: (selected: bool, product_id: int | None)} from the review form."""
    imp = PriceImport.objects.select_for_update().get(pk=imp.pk)
    _check_draft(imp)
    rows = {r.pk: r for r in imp.rows.select_related("product")}
    products = Product.objects.in_bulk([pid for _, pid in choices.values() if pid])
    chosen = []
    for row_id, (selected, product_id) in choices.items():
        row = rows.get(row_id)
        if row is None:
            continue
        product = products.get(product_id) if product_id else None
        if product_id and product is None:
            raise ValidationError(f"Рядок {row.row_no}: продукт не знайдено.")
        if product is not None and product != row.product:
            row.product, row.match = product, PriceImportRow.Match.MANUAL
            row.old_price = product.price_per_kg
        row.selected = bool(selected)
        if row.selected:
            if row.problem:
                raise ValidationError(f"Рядок {row.row_no}: {row.problem}")
            if row.product is None:
                raise ValidationError(f"Рядок {row.row_no} «{row.raw_name}»: оберіть продукт або зніміть позначку.")
            chosen.append(row)
    by_product = {}
    for row in chosen:
        other = by_product.setdefault(row.product.pk, row)
        if other is not row and other.price_per_kg != row.price_per_kg:
            raise ValidationError(f"«{row.product.name}» обрано в рядках {other.row_no} і {row.row_no} з різними цінами — залиште один.")
    changed = 0
    for row in chosen:
        product = row.product
        if product.price_per_kg != row.price_per_kg:
            product.price_per_kg = row.price_per_kg
            product._change_reason = f"Імпорт цін: {imp.filename}"
            product.save()
            changed += 1
        row.applied = True
        if row.match in (PriceImportRow.Match.MANUAL, PriceImportRow.Match.FUZZY):
            ProductAlias.objects.update_or_create(name=normalize(row.raw_name)[:300],
                                                  defaults={"product": product, "created_by": user})
    PriceImportRow.objects.bulk_update(list(rows.values()), ["product", "match", "old_price", "selected", "applied"])
    imp.status, imp.applied_at, imp.applied_by, imp.applied_count = PriceImport.Status.APPLIED, timezone.now(), user, changed
    imp.save(update_fields=["status", "applied_at", "applied_by", "applied_count"])
    return changed


def cancel(imp):
    _check_draft(imp)
    imp.status = PriceImport.Status.CANCELLED
    imp.save(update_fields=["status"])
