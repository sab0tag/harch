"""S4-02 / S4-03: consumption for a period and the accountant's write-off."""
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.catalog.models import Product
from apps.consumption import services
from apps.consumption.models import WriteOff
from apps.consumption.tests.conftest import DAY, as_user

pytestmark = pytest.mark.django_db
WO = f"/write-off/{DAY.isoformat()}/"


# ------------------------------------------------------------------ S4-02
def test_period_totals_equal_snapshot(cooking_day):
    data = services.period_consumption(DAY, DAY, "product")
    snap = services.snapshot_by_product(cooking_day.plan)
    assert data["brutto_g"] == sum(b for _, b in snap.values())
    assert {r["product_id"]: r["brutto_g"] for r in data["groups"][0]["rows"]} == {p: b for p, (_, b) in snap.items()}


def test_group_by_diet_sums_to_same_total(cooking_day):
    by_diet = services.period_consumption(DAY, DAY, "diet")
    assert {g["label"] for g in by_diet["groups"]} == {"ОВД", "ЩД", "№9"}
    assert by_diet["brutto_g"] == services.period_consumption(DAY, DAY)["brutto_g"]


def test_consumption_page_and_xlsx(client, cooking_day):
    c = as_user(client, cooking_day.accountant)
    q = f"?from={DAY.isoformat()}&to={DAY.isoformat()}"
    assert "РАЗОМ" in c.get("/consumption/" + q).content.decode()
    r = c.get("/consumption/" + q + "&format=xlsx")
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(r.content)).active
    names = [ws.cell(row=i, column=1).value for i in range(6, ws.max_row + 1)]
    assert "РАЗОМ" in names
    total_row = names.index("РАЗОМ") + 6
    assert ws.cell(row=total_row, column=2).value == pytest.approx(
        float(services.period_consumption(DAY, DAY)["brutto_g"] / 1000), abs=0.001)


def test_consumption_access(client, cooking_day):
    assert as_user(client, cooking_day.nurse["Терапевтичне"]).get("/consumption/").status_code == 403
    assert as_user(client, cooking_day.chef).get("/consumption/").status_code == 200


# ------------------------------------------------------------------ S4-03
def test_write_off_waits_for_kitchen(client, cooking_day):
    html = as_user(client, cooking_day.accountant).get(WO).content.decode()
    assert "Кухня ще не закрила день" in html


def test_missing_prices_block_submit_then_entered_on_page(client, write_off_day):
    w = write_off_day
    c = as_user(client, w.accountant)
    html = c.get(WO).content.decode()
    assert "Немає ціни" in html
    r = c.post(WO, {"action": "submit"}, follow=True)
    assert "Немає ціни за кг" in r.content.decode()
    w.write_off.refresh_from_db()
    assert w.write_off.status != "submitted"
    ids = [r["product_id"] for r in services.live_write_off_rows(w.plan)]
    c.post(WO, {"action": "prices", **{f"price_{pid}": "40,50" for pid in ids}})
    assert set(Product.objects.filter(pk__in=ids).values_list("price_per_kg", flat=True)) == {Decimal("40.50")}
    c.post(WO, {"action": "submit"})
    w.write_off.refresh_from_db()
    assert w.write_off.status == "submitted"


@pytest.fixture
def priced(write_off_day):
    Product.objects.update(price_per_kg=Decimal("100.00"))
    return write_off_day


def test_submit_freezes_lines_with_prices(client, priced):
    w = priced
    as_user(client, w.accountant).post(WO, {"action": "submit"})
    wo = WriteOff.objects.get(plan=w.plan)
    total_g = sum(l.quantity_g for l in wo.lines.all())
    assert wo.total_cost == (total_g / 1000 * 100).quantize(Decimal("0.01"))
    Product.objects.update(price_per_kg=Decimal("1.00"))   # later price change
    rows = services.write_off_rows(WriteOff.objects.get(plan=w.plan))
    assert all(r["price_per_kg"] == Decimal("100.00") for r in rows)


def test_quantities_come_from_kitchen_fact(client, priced):
    w = priced
    a = w.plan.actuals.first()
    a.actual_brutto_g = a.planned_brutto_g + 1000
    a.save()
    row = next(r for r in services.live_write_off_rows(w.plan) if r["product_id"] == a.product_id)
    assert row["quantity_g"] == a.actual_brutto_g


def test_dietician_approves_and_day_is_closed(client, priced):
    w = priced
    as_user(client, w.accountant).post(WO, {"action": "submit"})
    assert "Затвердити списання" not in as_user(client, w.accountant).get(WO).content.decode()
    as_user(client, w.dietician).post(f"/tasks/{w.write_off.pk}/approve/")
    wo = WriteOff.objects.get(plan=w.plan)
    assert wo.approved_at is not None
    html = as_user(client, w.accountant).get(WO).content.decode()
    assert "День закрито" in html


def test_write_off_access(client, priced):
    assert as_user(client, priced.chef).get(WO).status_code == 403
    assert as_user(client, priced.nurse["Терапевтичне"]).get(WO).status_code == 403
    html = as_user(client, priced.dietician).get(WO).content.decode()
    assert "Подати на звірку" not in html   # the accountant is the maker


def test_dietician_cannot_change_prices_via_write_off_before_it_is_hers(client, priced):
    pid = services.live_write_off_rows(priced.plan)[0]["product_id"]
    as_user(client, priced.dietician).post(WO, {"action": "prices", f"price_{pid}": "1"})
    assert Product.objects.get(pk=pid).price_per_kg == Decimal("100.00")


def test_undated_urls_post_to_the_same_day(client, priced):
    """/write-off/ and /kitchen/ without a date must not break on POST (found in browser walk)."""
    html = as_user(client, priced.accountant).get("/write-off/").content.decode()
    assert f'action="/write-off/{DAY.isoformat()}/"' in html
    r = as_user(client, priced.accountant).post("/write-off/", {"action": "submit"})
    assert r.status_code == 302


def test_price_import_does_not_touch_submitted_write_off(client, priced):
    """S5-07: a later price import changes the catalogue, never a write-off already handed in."""
    from apps.catalog import price_import as pi
    from apps.catalog.models import PriceImport, PriceImportRow

    w = priced
    as_user(client, w.accountant).post(WO, {"action": "submit"})
    product = WriteOff.objects.get(plan=w.plan).lines.first().product
    imp = PriceImport.objects.create(filename="new.xlsx", sheets={"s": [["Назва", "Ціна"], [product.name, "999"]]})
    pi.set_mapping(imp, "s", 0, 0, 1)
    row = imp.rows.get()
    assert row.match == PriceImportRow.Match.EXACT
    pi.apply(imp, w.accountant, {row.pk: (True, product.pk)})
    assert Product.objects.get(pk=product.pk).price_per_kg == Decimal("999.00")
    rows = services.write_off_rows(WriteOff.objects.get(plan=w.plan))
    assert all(r["price_per_kg"] == Decimal("100.00") for r in rows)


def test_missing_prices_round_trip_through_excel(client, write_off_day):
    """S5-07: export «без ціни» → accountant fills column D → import reads it back without asking."""
    from django.core.files.uploadedfile import SimpleUploadedFile
    from openpyxl import load_workbook

    from apps.catalog.models import PriceImport

    w = write_off_day
    c = as_user(client, w.accountant)
    missing = [r for r in services.live_write_off_rows(w.plan) if r["price_per_kg"] is None]
    assert missing
    html = c.get(WO).content.decode()
    assert "?export=missing-prices" in html and "/catalog/products/prices/import/?back=" in html

    r = c.get(WO + "?export=missing-prices")
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert [ws.cell(row=4, column=i).value for i in range(1, 6)] == ["№", "Найменування", "Од. вим.", "Ціна, грн", "До списання, кг"]
    names = [ws.cell(row=i, column=2).value for i in range(5, ws.max_row + 1)]
    assert sorted(names) == sorted(Product.objects.get(pk=m["product_id"]).name for m in missing)
    for i in range(5, ws.max_row + 1):
        ws.cell(row=i, column=4, value="87,50")                 # the accountant types prices, comma and all
    buf = BytesIO()
    wb.save(buf)

    c.post("/catalog/products/prices/import/", {"back": WO, "file": SimpleUploadedFile("filled.xlsx", buf.getvalue())})
    imp = PriceImport.objects.latest("pk")
    assert (imp.header_row, imp.name_col, imp.unit_col, imp.price_col) == (3, 1, 2, 3)   # recognised without asking
    rows = list(imp.rows.all())
    assert len(rows) == len(missing) and all(r.match == "exact" and r.selected for r in rows)
    c.post(f"/catalog/products/prices/import/{imp.pk}/",
           {"action": "apply", **{f"apply_{r.pk}": "1" for r in rows}, **{f"product_{r.pk}": r.product_id for r in rows}})
    done = c.get(f"/catalog/products/prices/import/{imp.pk}/").content.decode()
    assert f'href="{WO}"' in done and "До списання" in done
    assert all(r["price_per_kg"] == Decimal("87.50") for r in services.live_write_off_rows(w.plan))
    assert "Немає ціни" not in c.get(WO).content.decode()


def test_missing_prices_export_access_and_back_link_is_local(client, write_off_day):
    w = write_off_day
    assert as_user(client, w.chef).get(WO + "?export=missing-prices").status_code == 403
    html = as_user(client, w.accountant).get("/catalog/products/prices/import/?back=https://evil.example/x").content.decode()
    assert "evil.example" not in html


def test_a_price_changed_in_a_write_off_is_that_days_only(client, priced):
    """A correction of 25.09's price must not change the catalog (today's price) or other days."""
    w = priced
    pid = services.live_write_off_rows(w.plan)[0]["product_id"]
    c = as_user(client, w.accountant)
    c.post(WO, {"action": "prices", f"price_{pid}": "120"})
    assert Product.objects.get(pk=pid).price_per_kg == Decimal("100.00")            # catalog untouched
    row = next(r for r in services.live_write_off_rows(w.plan) if r["product_id"] == pid)
    assert row["price_per_kg"] == Decimal("120") and row["day_price"]
    assert "ціна дня" in c.get(WO).content.decode()
    c.post(WO, {"action": "submit"})
    line = WriteOff.objects.get(plan=w.plan).lines.get(product_id=pid)
    assert line.price_per_kg == Decimal("120.00")                                   # frozen with the day's price


def test_setting_the_catalog_price_back_drops_the_day_price(client, priced):
    from apps.consumption.models import WriteOffPrice

    pid = services.live_write_off_rows(priced.plan)[0]["product_id"]
    c = as_user(client, priced.accountant)
    c.post(WO, {"action": "prices", f"price_{pid}": "120"})
    c.post(WO, {"action": "prices", f"price_{pid}": "100"})
    assert not WriteOffPrice.objects.exists()
