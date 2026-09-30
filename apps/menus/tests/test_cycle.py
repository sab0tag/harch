"""S7-04: the 7-day menu cycle from Excel, and the day plan picking the menu by weekday."""
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Role
from apps.catalog.models import Diet, Dish
from apps.menus import cycle
from apps.menus.models import Menu
from apps.planning.services import suggest_menu

pytestmark = pytest.mark.django_db
DAYS = ["Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця", "Субота", "Неділя"]


def workbook(rows, header=("Дієта", "День", "Прийом їжі", "Страва (№ картки або назва)")):
    wb = Workbook()
    ws = wb.active
    ws.title = "Меню"
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "menu.xlsx"
    return buf


def dishes_for(code, n=2):
    return list(Dish.objects.filter(diets__code=code, is_active=True).order_by("card_no")[:n])


def week_rows(code):
    a, b = dishes_for(code)
    return [r for day in DAYS for r in ((code, day, "Обід", a.card_no), (code, day, "Вечеря", b.name))]


def test_template_has_every_diet_day_and_meal_and_the_cards(seeded):
    wb = load_workbook(BytesIO(cycle.template_xlsx().content))
    assert wb.sheetnames == ["Меню", "Картки"]
    diets = cycle.cycle_diets().count()
    assert wb["Меню"].max_row == 1 + diets * 7 * 4                   # сніданок, 2-й сніданок, обід, вечеря
    assert wb["Картки"].max_row == 1 + Dish.objects.filter(is_active=True).count()


def test_import_builds_the_week_and_the_plan_picks_by_weekday(seeded):
    rows, errors = cycle.parse(workbook(week_rows("ОВД") + [("ЩД", "Пн", "обід", dishes_for("ЩД")[0].card_no)]))
    assert errors == []
    week = cycle.preview(rows)
    by_code = {d.code: e for d, e in week.items()}
    assert by_code["ОВД"]["missing"] == [] and len(by_code["ЩД"]["missing"]) == 6
    old_ovd = set(Menu.objects.filter(diet__code="ОВД", is_active=True).values_list("pk", flat=True))
    created = cycle.apply(rows)
    assert len(created) == 8
    ovd = Menu.objects.filter(diet__code="ОВД", is_active=True)
    assert sorted(ovd.values_list("cycle_day", flat=True)) == [1, 2, 3, 4, 5, 6, 7]
    assert not set(ovd.values_list("pk", flat=True)) & old_ovd       # the generated menus are switched off
    thursday = date(2026, 10, 1)
    assert suggest_menu(Diet.objects.get(code="ОВД"), thursday).cycle_day == 4
    assert suggest_menu(Diet.objects.get(code="ЩД"), thursday) is None   # no Thursday in ЩД's cycle yet
    assert [i.meal_slot for i in ovd.get(cycle_day=1).items.all()] == ["lunch", "dinner"]


def test_reimport_replaces_the_previous_cycle(seeded):
    cycle.apply(cycle.parse(workbook(week_rows("ОВД")))[0])
    first = set(Menu.objects.filter(diet__code="ОВД", is_active=True).values_list("pk", flat=True))
    cycle.apply(cycle.parse(workbook(week_rows("ОВД")))[0], deactivate_others=False)
    second = set(Menu.objects.filter(diet__code="ОВД", is_active=True).values_list("pk", flat=True))
    assert len(second) == 7 and not first & second


@pytest.mark.parametrize("row, message", [
    (("XYZ", "Пн", "Обід", 1), "дієта «XYZ»"),
    (("ОВД", "Свято", "Обід", 1), "день «Свято»"),
    (("ОВД", "Пн", "Полудень", 1), "прийом їжі «Полудень»"),
    (("ОВД", "Пн", "Обід", "Страва-привид"), "немає такої картки"),
])
def test_one_bad_row_changes_nothing(seeded, row, message):
    before = Menu.objects.filter(is_active=True).count()
    rows, errors = cycle.parse(workbook(week_rows("ОВД") + [row]))
    assert any(message in e and e.startswith("Рядок 16") for e in errors)
    assert Menu.objects.filter(is_active=True).count() == before


def test_a_card_gets_the_diet_allowed(seeded):
    dish = Dish.objects.exclude(diets__code="НКД").filter(is_active=True).first()
    rows, _ = cycle.parse(workbook([("НКД", "1", "Обід", dish.card_no)]))
    assert [d.card_no for e in cycle.preview(rows).values() for d in e["cards_to_allow"]] == [dish.card_no]
    cycle.apply(rows)
    assert dish.diets.filter(code="НКД").exists()


def test_the_dietician_imports_through_the_screens(client, login, seeded):
    c = login(Role.DIETICIAN)
    assert c.get("/menus/cycle/template.xlsx").status_code == 200
    r = c.post("/menus/cycle/", {"file": workbook(week_rows("ОВД"))})
    assert r.status_code == 302 and r.url == "/menus/cycle/preview/"
    assert "Понеділок" in c.get(r.url).content.decode()
    r = c.post("/menus/cycle/preview/", {"deactivate_others": "on"}, follow=True)
    assert "Створено меню 7-денного циклу: 7" in r.content.decode()
    bad = c.post("/menus/cycle/", {"file": workbook([("ОВД", "Пн", "Обід", "Страва-привид")])})
    assert "Нічого не змінено" in bad.content.decode()


@pytest.mark.parametrize("role", [Role.CHEF, Role.ACCOUNTANT, Role.STOREKEEPER])
def test_only_reference_editors_import(client, login, seeded, role):
    assert login(role).get("/menus/cycle/").status_code == 403
