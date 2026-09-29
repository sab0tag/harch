"""S6-05: pilot staff from a spreadsheet — create or update accounts, hand out temporary passwords.

Input: CSV (`;` or `,`, UTF-8) or .xlsx with a header row. Columns, any order, Ukrainian or English:
логін (optional — generated from the surname), ПІБ, роль, відділення (required for a head nurse), телефон.
Every row is checked before anything is written: one bad row → nothing changes.
"""
import csv
import re
import secrets
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction

from apps.orgstructure.models import Department

from .models import Role, User

HEADERS = {
    "login": ("логін", "логин", "login", "username"),
    "name": ("піб", "пiб", "прізвище ім'я по батькові", "прізвище, ім'я, по батькові", "фио", "name"),
    "role": ("роль", "посада", "role"),
    "department": ("відділення", "отделение", "department"),
    "phone": ("телефон", "phone"),
}

ROLE_ALIASES = {
    Role.ADMIN: ("адміністратор", "адмін"),
    Role.DIETICIAN: ("дієтсестра", "лікар-дієтолог", "дієтолог", "сестра медична з дієтичного харчування"),
    Role.CHEF: ("завідувач виробництва", "шеф-кухар", "шеф", "кухар"),
    Role.DEPT_NURSE: ("старша медсестра", "старша медична сестра", "медсестра"),
    Role.ACCOUNTANT: ("бухгалтер",),
}

# Official Ukrainian romanization (КМУ №55, 2010): є ї й ю я are written differently at the start of a word.
TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "h", "ґ": "g", "д": "d", "е": "e", "є": "ie", "ж": "zh", "з": "z",
    "и": "y", "і": "i", "ї": "i", "й": "i", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p",
    "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh",
    "щ": "shch", "ь": "", "ю": "iu", "я": "ia", "'": "", "’": "", "ʼ": "",
}
WORD_START = {"є": "ye", "ї": "yi", "й": "y", "ю": "yu", "я": "ya"}
PASSWORD_ALPHABET = "abcdefghjkmnpqrstuvwxyz23456789"  # no 0/o, 1/l/i: read aloud or copied by hand


def translit(text):
    text = text.lower().replace("зг", "zgh")
    out = []
    for i, ch in enumerate(text):
        start = i == 0 or not (text[i - 1].isalpha() or text[i - 1] in "'’ʼ")  # Мар'яна → mariana
        out.append(WORD_START[ch] if start and ch in WORD_START else TRANSLIT.get(ch, ch))
    return re.sub(r"[^a-z0-9]", "", "".join(out))


def temporary_password():
    pick = lambda n: "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(n))  # noqa: E731
    return f"{pick(4)}-{pick(4)}-{pick(2)}"


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("’", "'").replace("ʼ", "'"))


def _role(value):
    v = _norm(value)
    for role, aliases in ROLE_ALIASES.items():
        if v in (role.value, _norm(role.label), *aliases):
            return role
    return None


def read_table(path):
    """Rows as dicts keyed by login/name/role/department/phone; the first non-empty row is the header."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        rows = [list(r) for r in load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,\t")
        rows = list(csv.reader(text.splitlines(), dialect))
    rows = [r for r in rows if any(_norm(c) for c in r)]
    if not rows:
        return []
    columns = {}
    for i, title in enumerate(rows[0]):
        for key, names in HEADERS.items():
            if _norm(title) in names:
                columns[key] = i
    missing = {"name", "role"} - columns.keys()
    if missing:
        raise ValueError("У таблиці немає колонок: " + ", ".join(HEADERS[k][0] for k in sorted(missing)))
    return [
        {"line": n, **{k: str(r[i]).strip() if i < len(r) and r[i] is not None else "" for k, i in columns.items()}}
        for n, r in enumerate(rows[1:], start=2)
    ]


@dataclass
class Result:
    created: list = field(default_factory=list)
    updated: list = field(default_factory=list)
    deactivated: list = field(default_factory=list)
    errors: list = field(default_factory=list)


def import_users(rows, *, reset_passwords=False, deactivate_others=False, dry_run=False):
    """Validate every row, then create/update in one transaction. New or reset accounts get a temporary
    password (shown once, in the result) and must change it at the first sign-in."""
    result = Result()
    departments = {_norm(d.name): d for d in Department.objects.filter(is_active=True)}
    taken = set(User.objects.values_list("username", flat=True))
    planned = []
    for row in rows:
        where = f"рядок {row['line']}"
        parts = row.get("name", "").split()
        role = _role(row.get("role"))
        if len(parts) < 2:
            result.errors.append(f"{where}: ПІБ — щонайменше прізвище та ім'я")
        if role is None:
            result.errors.append(f"{where}: невідома роль «{row.get('role')}»")
        department = None
        if role == Role.DEPT_NURSE:
            name = _norm(re.sub(r"(?i)\s*відділення\s*", " ", row.get("department", "")))
            department = departments.get(name)
            if department is None:
                result.errors.append(f"{where}: відділення «{row.get('department')}» не знайдено "
                                     f"(є: {', '.join(sorted(d.name for d in departments.values()))})")
        if result.errors and result.errors[-1].startswith(where):
            continue
        login = row.get("login", "").strip().lower()
        if not login:
            base = translit(parts[0]) + "." + translit(parts[1])[:1]
            login, n = base, 2
            while login in taken:
                login, n = f"{base}{n}", n + 1
        elif not re.fullmatch(r"[a-z0-9._@+-]+", login):
            result.errors.append(f"{where}: логін «{login}» — лише латиниця, цифри та . _ -")
            continue
        taken.add(login)
        planned.append({"login": login, "last_name": parts[0], "first_name": " ".join(parts[1:]),
                        "role": role, "department": department, "phone": row.get("phone", "")})
    logins = [p["login"] for p in planned]
    for dup in sorted({x for x in logins if logins.count(x) > 1}):
        result.errors.append(f"логін «{dup}» повторюється у таблиці")
    if result.errors:
        return result

    with transaction.atomic():
        for p in planned:
            user = User.objects.filter(username=p["login"]).first()
            fields = {"last_name": p["last_name"], "first_name": p["first_name"], "role": p["role"],
                      "department": p["department"], "phone": p["phone"], "is_active": True}
            password = None
            if user is None:
                user = User(username=p["login"], **fields)
                password = temporary_password()
                bucket = result.created
            else:
                for k, v in fields.items():
                    setattr(user, k, v)
                if reset_passwords:
                    password = temporary_password()
                bucket = result.updated
            if password:
                user.set_password(password)
                user.must_change_password = True
            if not dry_run:
                user.save()
            bucket.append({"login": p["login"], "name": f"{p['last_name']} {p['first_name']}",
                           "role": Role(p["role"]).label, "department": p["department"].name if p["department"] else "",
                           "password": password or ""})
        if deactivate_others:
            others = User.objects.filter(is_active=True, is_superuser=False).exclude(username__in=logins)
            result.deactivated = sorted(others.values_list("username", flat=True))
            if not dry_run:
                others.update(is_active=False)
        if dry_run:
            transaction.set_rollback(True)
    return result
